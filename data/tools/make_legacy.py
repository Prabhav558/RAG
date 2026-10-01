"""Build the legacy-spreadsheet datasets used to exercise the Cycle 2 migration (data/legacy/).

Each file mimics how teams actually keep scorecards in Excel/CSV today, with deliberately injected problems.
The expected outcome of every injected problem is tested in backend/tests/test_migration.py.

    python data/tools/make_legacy.py
"""

from __future__ import annotations

import csv
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

OUT = Path(__file__).resolve().parents[2] / "data" / "legacy"
rng = random.Random(7)


def weighted(scores: dict, weights: dict) -> float:
    return round(sum(scores[k] * weights[k] for k in scores) / sum(weights[k] for k in scores), 2)


# ---------------------------------------------------------------- L1: clean, flat, 1-5
def code_review():
    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard"
    for row in [["Scorecard: Code Review Quality"], ["Purpose", "Catch defects and design problems before merge"],
                ["Objective", "Every reviewed PR scores 4 or above"], ["Scale", "1-5"], ["Target", 4],
                ["Subject", "Task"], []]:
        ws.append(row)
    ws.append(["No", "KPI", "Weight %", "Description", "5", "4", "3", "2", "1"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    kpis = [
        ("1", "Defects found", 30, "Real defects caught", "All real defects caught with repro", "Most defects caught",
         "Obvious defects only", "Missed a defect", "Rubber stamp"),
        ("2", "Design feedback", 25, "Architecture and design", "Improves design materially", "Useful design notes",
         "Some design comments", "Style only", "None"),
        ("3", "Clarity of comments", 20, "Actionable wording", "Every comment actionable", "Mostly actionable",
         "Some vague comments", "Mostly vague", "Unclear"),
        ("4", "Turnaround", 15, "Time to first review", "< 4 hours", "< 1 day", "< 2 days", "< 4 days", ">= 4 days"),
        ("5", "Tone", 10, "Respectful, constructive", "Exemplary", "Constructive", "Neutral", "Curt", "Hostile"),
    ]
    for k in kpis:
        ws.append(list(k))
    w = {k[0]: k[2] for k in kpis}
    rs = wb.create_sheet("Ratings")
    rs.append(["PR", "Reviewer", "Date", "Defects found", "Design feedback", "Clarity of comments", "Turnaround",
               "Tone", "Total", "Comments"])
    d0 = date(2026, 3, 2)
    for i in range(15):
        s = {k: rng.randint(2, 5) for k in w}
        rs.append([f"PR-{1200 + i}", rng.choice(["Asha", "Ben", "Chen"]), d0 + timedelta(days=3 * i),
                   *s.values(), weighted(s, w), rng.choice(["", "Good catch on the race condition", "Needs follow-up"])])
    wb.save(OUT / "code-review-quality.xlsx")


# ---------------------------------------------------------------- L2: messy, hierarchical, 0-10, M01-M05
def training_feedback():
    wb = Workbook()
    ws = wb.active
    ws.title = "KPIs"
    for row in [["TRAINING SESSION FEEDBACK SCORECARD"], ["Scorecard: Training Session Quality"],
                ["Owner: L&D team"], ["Target: good"], ["Last updated", "Q4 2025"], []]:
        ws.append(row)
    ws.append(["S.No", "Parameter", "Weightage", "Score 10", "Score 9", "Score 8", "Score 6-7", "Score 4-5", "Score 0-3",
               "Owner notes"])
    rows = [
        ("1", "Content", 40, None, None, None, None, None, None, ""),
        ("1.1", "Relevance to role", 60, "Directly applicable", "Directly applicable", "Mostly applicable",
         "Partly applicable", "Rarely applicable", "Irrelevant", ""),
        ("1.2", "Accuracy", "high", "No errors", "No errors", "One minor slip", "Several slips", "Errors", "Wrong", ""),
        ("2", "Delivery", 50, None, None, None, None, None, None, ""),
        ("2.1", "Clarity", 40, "Crystal clear", "Very clear", "Clear", "Some confusion", None, "Confusing", "check 4-5"),
        ("2.2", "Engagement", None, "Everyone engaged", "Most engaged", "Engaged", "Mixed", "Low", "None", ""),
        ("2.3", "Pace", 20, "Perfect", "Good", "Good", "Uneven", "Too fast/slow", "Unusable", ""),
        ("3", "Logistics", 20, None, None, None, None, None, None, ""),
        ("3.1", "Venue & equipment", 100, "Flawless", "Minor issue", "Minor issue", "Issues", "Major issues", "Failed", ""),
        ("", "", None, None, None, None, None, None, None, ""),
        ("", "Total", 110, None, None, None, None, None, None, ""),
    ]
    for r in rows:
        ws.append(list(r))
    # legacy formula: sum(weight * score) / 100 over the leaves, using the sheet's own (unnormalised) weights
    leaf_w = {"Relevance to role": 60 * 0.4, "Accuracy": 3 * 0.4, "Clarity": 40 * 0.5, "Engagement": 1 * 0.5,
              "Pace": 20 * 0.5, "Venue & equipment": 100 * 0.2}
    rs = wb.create_sheet("Feedback log")
    rs.append(["Session", "Assessor", "Date", *leaf_w.keys(), "Overall", "Remarks"])

    def add(subject, assessor, when, scores, overall=None, remarks=""):
        num = {k: v for k, v in scores.items() if isinstance(v, (int, float))}
        if overall is None and len(num) == len(scores):
            overall = round(sum(v * leaf_w[k] for k, v in num.items()) / 100, 2)
        rs.append([subject, assessor, when, *scores.values(), overall, remarks])

    d0 = date(2025, 10, 6)
    for i in range(14):
        add(f"Onboarding batch {i + 1}", rng.choice(["R. Iyer", "M. Das"]), d0 + timedelta(days=7 * i),
            {k: rng.randint(5, 10) for k in leaf_w})
    full = {k: 8 for k in leaf_w}
    add("Workshop: Git", "R. Iyer", "2025-12-01", {**full, "Venue & equipment": "N/A"}, remarks="Remote session")
    add("Workshop: SQL", "R. Iyer", "2025-12-02", {**full, "Pace": 7.5})
    add("Workshop: Cloud", "M. Das", "2025-12-03", {**full, "Clarity": 12})
    add("Workshop: Testing", "M. Das", "2025-12-04", {**full, "Engagement": "good"})
    add("Workshop: Docker", "M. Das", "2025-12-05", {**full, "Accuracy": None})
    add("Workshop: APIs", None, "2025-12-06", full)
    add("Workshop: Security", "R. Iyer", "31/02/2025", full)
    add("Workshop: Python", "R. Iyer", "2025-12-08", full)
    add("Workshop: Python", "R. Iyer", "2025-12-08", full)  # exact duplicate
    add("Workshop: Agile", "M. Das", "2025-12-09", full)
    add("Workshop: Agile", "M. Das", "2025-12-09", {**full, "Clarity": 5})  # conflicting duplicate
    add("Workshop: Kafka", "M. Das", "2025-12-10", {**full, "Clarity": 5})
    add("Workshop: Kafka", "M. Das", "2026-01-12", full)  # re-run: second attempt
    add("", "M. Das", "2025-12-11", full)  # no subject
    add("Workshop: Linux", "R. Iyer", "2025-12-12", {k: None for k in leaf_w})  # no scores
    rs.append([])
    rs.append(["Average", None, None, *[7.9] * len(leaf_w), 7.9, ""])
    wb.save(OUT / "training-session-quality.xlsx")


# ---------------------------------------------------------------- L3: CSV, indentation, %, no guidelines
def vendor_csv():
    with open(OUT / "vendor-assessment.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Scorecard: Vendor Assessment"])
        w.writerow(["Scale: %"])
        w.writerow(["Target: 75"])
        w.writerow(["Applies to: Vendor"])
        w.writerow([])
        w.writerow(["Criteria", "Weight", "Critical"])
        for name, wt, crit in [("Commercials", 30, ""), ("  Pricing competitiveness", 60, ""),
                               ("  Payment terms", 40, ""), ("Delivery", 40, ""), ("  On-time delivery", 70, "Y"),
                               ("  Quality of goods", 30, "Y"), ("Compliance", 30, ""), ("  ISO certification", 50, ""),
                               ("  Data protection", 50, "yes")]:
            w.writerow([name, wt, crit])
    with open(OUT / "vendor-assessment-ratings.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Vendor", "Rated by", "Date", "Pricing competitiveness", "Payment terms", "On-time delivery",
                    "Quality of goods", "ISO certification", "Data protection", "Score"])
        for i, v in enumerate(["Acme Supplies", "Globex", "Initech", "Umbrella Corp", "Hooli", "Stark Industries"]):
            s = [rng.randint(45, 100) for _ in range(6)]
            w.writerow([v, "Procurement", f"0{i + 1}.02.2026", *s, ""])


# ---------------------------------------------------------------- L4: garbage
def broken():
    wb = Workbook()
    ws = wb.active
    ws.append(["Quarterly numbers"])
    ws.append(["Region", "Revenue", "Growth"])
    ws.append(["North", 120, 0.05])
    wb.save(OUT / "broken-no-header.xlsx")
    (OUT / "corrupt.xlsx").write_bytes(b"PK\x03\x04 this is not really a zip")


# ---------------------------------------------------------------- L5: unsupported scale (needs rescale)
def seven_point():
    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard"
    for row in [["Scorecard: Meeting Effectiveness"], ["Scale: 1 to 7"], ["Target: 5"], ["Subject: Meeting"], []]:
        ws.append(row)
    ws.append(["KPI", "Weight", "7", "6", "5", "4", "3", "2", "1"])
    for name, wt in [("Agenda shared in advance", 1), ("Decisions recorded", 2), ("Actions with owners", 2),
                     ("Timekeeping", 1)]:
        ws.append([name, wt, "Always", "Almost always", "Usually", "Sometimes", "Rarely", "Almost never", "Never"])
    ws.cell(row=7, column=1).alignment = Alignment(indent=0)
    rs = wb.create_sheet("Ratings")
    rs.append(["Meeting", "Rated by", "Date", "Agenda shared in advance", "Decisions recorded", "Actions with owners",
               "Timekeeping"])
    for i in range(8):
        rs.append([f"Weekly sync {i + 1}", "PMO", date(2026, 1, 5) + timedelta(days=7 * i),
                   *[rng.randint(3, 7) for _ in range(4)]])
    wb.save(OUT / "meeting-effectiveness-7pt.xlsx")


# ---------------------------------------------------------------- L6: a clean, well-formed sheet (the rule-based importer reads it as-is)
def client_email():
    """Client Email Quality in the layout the rule-based importer recognises: metadata lines, dotted numbering,
    weights, one guideline column per score range, a critical flag, and a ratings sheet whose legacy totals the
    new engine reproduces (the totals use the sheet's own weights, nested like the hierarchy)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard"
    for row in [["Scorecard: Client Email Quality"],
                ["Purpose", "Make sure every client email is clear, accurate and professional before it is sent"],
                ["Scope", "Subject line, body, tone and attachments of external client emails. Internal threads are out."],
                ["Objective", "An email is ready to send at 7 or above overall, with no spelling or fact errors"],
                ["Scale", "0-10"], ["Target", 7], ["Subject", "Communication"], ["Owner", "Sales operations"], []]:
        ws.append(row)
    heads = ["No", "KPI", "Weight %", "Description", "Critical", "9-10", "8", "7", "6", "5", "4", "0-3"]
    ws.append(heads)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    leaves = [
        ("1.1", "Purpose stated", 50, "Reader knows why the email was sent", "",
         "Purpose in the subject line and the first sentence", "Purpose in the first 2 sentences, subject names the topic",
         "Purpose in the first 2 sentences, subject is generic", "Purpose only later in the first paragraph",
         "Purpose must be inferred", "Purpose unclear to most readers", "No purpose"),
        ("1.2", "Ask is explicit", 50, "Reader knows what is wanted and by when", "",
         "One clear ask with a specific date or time", "Clear ask with a specific date or a date range",
         "Clear ask, deadline only vague (soon)", "Clear ask but no deadline, or 'asap'", "Ask buried among other points",
         "Ask unclear or contradictory", "No ask at all"),
        ("2.1", "Professional tone", 50, "Polite, confident, right for a client", "",
         "Greeting and sign-off, no informal phrases", "At most 1 informal phrase", "2 informal phrases, nothing negative",
         "3 or more informal phrases or no greeting", "Slang or curt phrasing in places",
         "Dismissive, sarcastic or defensive lines", "Rude or inflammatory"),
        ("2.2", "No errors", 50, "Spelling, grammar and facts are correct", "Y",
         "0 spelling or grammar errors, 0 factual errors", "1 minor error", "2 minor errors", "3-4 errors, meaning intact",
         "5-6 errors or one minor factual slip", "7 or more errors or a wrong name, date or figure",
         "Pervasive errors or a serious factual mistake"),
        ("3.1", "Owners and dates stated", 50, "Each action has an owner and a date", "",
         "100% of actions have an owner and a date", "At least 85% have both", "70-84% have both", "50-69% have both",
         "Under 50% have an owner", "Actions listed, none assigned", "No actions identifiable"),
        ("3.2", "Next steps explicit", 50, "Reader knows what happens next", "",
         "Next step stated for every open point", "One open point without a next step", "Next steps stated only generally",
         "2 or more open points without a next step", "Next steps must be inferred",
         "Decisions and open points mixed together", "Nothing actionable"),
    ]
    groups = [("1", "Clarity", 40), ("2", "Tone", 30), ("3", "Actionability", 30)]
    for g, name, w in groups:
        ws.append([g, name, w, "", ""])
        for leaf in [x for x in leaves if x[0].startswith(g + ".")]:
            ws.append(list(leaf))
    ws.append([])
    ws.append(["", "Total", 100])
    gw = {g: w for g, _, w in groups}
    leaf_w = {name: gw[code.split(".")[0]] / sum(gw.values()) * lw / sum(
        x[2] for x in leaves if x[0].split(".")[0] == code.split(".")[0]) for code, name, lw, *_ in leaves}
    rs = wb.create_sheet("Ratings")
    names = list(leaf_w)
    rs.append(["Subject", "Reviewer", "Date", *names, "Total", "Comments"])
    samples = [
        ("Phase 2 budget approval", [10, 10, 10, 10, 9, 9], "Model email"),
        ("Project status update", [8, 7, 9, 10, 7, 8], ""),
        ("Delivery date query", [7, 7, 8, 9, 6, 7], ""),
        ("Invoice reminder", [6, 6, 8, 9, 5, 6], "Ask has no date"),
        ("Holiday cover notice", [8, 6, 9, 10, 6, 7], ""),
        ("Complaint reply", [5, 6, 5, 7, 5, 6], "Too casual for a complaint"),
        ("Rushed follow-up", [4, 4, 4, 4, 3, 4], "Typos and no clear ask"),
        ("Angry escalation", [3, 5, 2, 3, 2, 3], "Should not have been sent"),
    ]
    d0 = date(2026, 8, 3)
    for i, (subject, sc, note) in enumerate(samples):
        total = round(sum(v * leaf_w[n] for v, n in zip(sc, names)), 2)
        rs.append([subject, ["Priya", "Arjun"][i % 2], d0 + timedelta(days=4 * i), *sc, total, note])
    wb.save(OUT / "client-email-quality.xlsx")


# ---------------------------------------------------------------- L7: vague KPIs, no guidelines (for the AI-assisted import)
def vendor_review_vague():
    """The kind of sheet teams really keep: five vague KPI names, weights as words, one-line hints instead of
    guidelines, an 'Overall' row and a Total. The rule-based importer would need placeholders for every
    guideline; the AI-assisted import is meant to turn this into a real scorecard."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendor review"
    for row in [["Vendor review sheet"], ["Fill in after each quarterly review. Rate 1-5."], []]:
        ws.append(row)
    ws.append(["What", "Importance", "Notes"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in [("Quality", "high", "are they good?"), ("Communication", "medium", "do they reply"),
                ("Price", "high", "fair?"), ("Reliability", "high", "on time and so on"),
                ("Support", "low", "after sale help"), ("Overall impression", "", "gut feel"), ("Total", "", "")]:
        ws.append(list(row))
    rs = wb.create_sheet("Scores")
    rs.append(["Vendor", "Quality", "Communication", "Price", "Reliability", "Support", "Overall impression", "Comments"])
    r = random.Random(11)
    for name in ["Northwind Supplies", "Contoso Logistics", "Fabrikam Parts", "Tailspin Packaging", "Litware Services"]:
        rs.append([name, *[r.randint(2, 5) for _ in range(6)], r.choice(["ok", "", "slow in March", "great team"])])
    wb.save(OUT / "vendor-review-vague.xlsx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    code_review()
    training_feedback()
    vendor_csv()
    broken()
    seven_point()
    client_email()
    vendor_review_vague()
    for p in sorted(OUT.iterdir()):
        print(p.relative_to(OUT.parent.parent), p.stat().st_size, "bytes")


if __name__ == "__main__":
    main()
